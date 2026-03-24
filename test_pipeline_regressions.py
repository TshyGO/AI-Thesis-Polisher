import unittest
import uuid
from pathlib import Path

import openpyxl

from engine.excel_exporter import ExcelExporter
from engine.pipelines import PolishingPipeline


class FakeLLMClient:
    def __init__(self):
        self.model = "fake-model"
        self.stage0_calls = []
        self.stage1_prompts = []
        self.stage2_prompts = []

    def call_api(self, messages, temperature=0.3, timeout=60, max_retries=3):
        self.stage0_calls.append(messages)
        return "本章研究 APTES 有机硅体系；APTES、PMSA 属于关键术语，不应误判。"

    def call_json_api(self, messages, temperature=0.3, timeout=60, max_retries=3):
        prompt = messages[-1]["content"]

        if "Stage 1 提名的修改建议" in prompt or "[Proposed Revisions]" in prompt:
            self.stage2_prompts.append(prompt)
            return [
                {
                    "sentence_id": "S1",
                    "old": "值得注意的是",
                    "new": "",
                    "reason": "删除空泛套话",
                    "priority": "高",
                }
            ]

        self.stage1_prompts.append(prompt)
        if "[S1] 值得注意的是" in prompt:
            return [
                {
                    "sentence_id": "S1",
                    "old": "值得注意的是",
                    "new": "",
                    "reason": "删除空泛套话",
                    "priority": "高",
                }
            ]
        return []


class FakeDocumentProcessor:
    def __init__(self):
        self.paragraphs = {
            1: "值得注意的是，本研究聚焦APTES和PMSA协同改性体系。",
            2: "样品在80 ℃下被加热2 h，然后进行力学测试。",
        }
        self.revisions = []
        self.save_calls = 0
        self.open_calls = []

    def open_document(self, doc_path: str, read_only=False, track_revisions=True):
        self.open_calls.append((doc_path, read_only, track_revisions))

    def get_total_paragraphs(self) -> int:
        return len(self.paragraphs)

    def parse_chapters(self) -> list:
        return [{"name": "第1章_测试/引言", "start": 1, "end": 2}]

    def get_paragraph_text(self, global_idx: int) -> str:
        return self.paragraphs.get(global_idx, "")

    def get_neighbor_text(self, center_idx: int, direction: str = "prev", limit: int = 2) -> str:
        indexes = []
        if direction == "prev":
            indexes = range(max(1, center_idx - limit), center_idx)
        else:
            indexes = range(center_idx + 1, min(len(self.paragraphs) + 1, center_idx + limit + 1))
        return "\n".join(self.paragraphs[idx] for idx in indexes if self.paragraphs.get(idx))

    def apply_tracked_revision(self, global_idx: int, old_text: str, new_text: str) -> bool:
        self.revisions.append((global_idx, old_text, new_text))
        return True

    def save(self):
        self.save_calls += 1


class PipelineRegressionTests(unittest.TestCase):
    def _workspace_tempdir(self):
        root = Path(__file__).resolve().parent / "cache" / "test_artifacts" / uuid.uuid4().hex
        root.mkdir(parents=True, exist_ok=True)
        return root

    def _build_pipeline(self, cache_root: Path):
        client = FakeLLMClient()
        doc_parser = FakeDocumentProcessor()
        pipeline = PolishingPipeline(
            llm_client=client,
            doc_parser=doc_parser,
            config={
                "language": "chinese",
                "intensity": "standard",
                "min_chars": 5,
                "use_cross_review": True,
                "original_filename": "demo.docx",
                "skipped_chapters": [],
                "output_dir": str(cache_root / "exports"),
                "excel_output_filename": "Report_demo.xlsx",
                "prompt_customization": {
                    "stage0": "请特别关注术语统一。",
                    "stage1": "优先清理空泛套话，但不要误伤实验过程描述。",
                    "stage2": "对被动句保持保守，除非明显欧化。",
                },
            },
        )
        pipeline.cache_dir = cache_root
        pipeline.cache_file = cache_root / "workflow_cache.json"
        pipeline.chapter_notes_dir = cache_root / "chapter_notes"
        pipeline.chapter_notes_dir.mkdir(parents=True, exist_ok=True)
        pipeline.save_cache({})
        return pipeline, client, doc_parser

    def test_pipeline_restores_stage0_excel_fields_and_deletion_write(self):
        temp_root = self._workspace_tempdir()
        doc_path = temp_root / "demo.docx"
        doc_path.write_bytes(b"placeholder")

        pipeline, client, doc_parser = self._build_pipeline(temp_root / "cache")
        changes, excel_path = pipeline.process_document(str(doc_path))

        self.assertEqual(changes, 1)
        self.assertEqual(len(client.stage0_calls), 1)
        self.assertIn("本章研究 APTES 有机硅体系", client.stage1_prompts[0])
        self.assertIn("本章研究 APTES 有机硅体系", client.stage2_prompts[0])
        self.assertIn("优先清理空泛套话", client.stage1_prompts[0])
        self.assertIn("对被动句保持保守", client.stage2_prompts[0])
        self.assertIn((1, "值得注意的是", ""), doc_parser.revisions)
        self.assertEqual(Path(excel_path).name, "Report_demo.xlsx")
        self.assertEqual(Path(excel_path).parent.name, "exports")

        workbook = openpyxl.load_workbook(excel_path)
        self.assertIn("第1章_测试_引言", workbook.sheetnames)
        sheet = workbook["第1章_测试_引言"]
        self.assertEqual(sheet["B2"].value, "S1")
        self.assertEqual(sheet["I2"].value, "高")
        self.assertEqual(sheet["H2"].value, "✅保留")
        self.assertNotIn("值得注意的是", sheet["G2"].value)
        self.assertIn("APTES", sheet["G2"].value)

    def test_excel_exporter_creates_chapter_sheets_even_without_records(self):
        temp_root = self._workspace_tempdir()
        export_path = temp_root / "report.xlsx"
        exporter = ExcelExporter()
        exporter.export(
            str(export_path),
            records=[],
            chapters=[
                {"name": "引言/前言", "start": 1, "end": 10},
                {"name": "第2章:方法", "start": 11, "end": 20},
            ],
        )

        workbook = openpyxl.load_workbook(export_path)
        self.assertIn("引言_前言", workbook.sheetnames)
        self.assertIn("第2章_方法", workbook.sheetnames)
        self.assertEqual(workbook["引言_前言"]["A1"].value, "段落号")


if __name__ == "__main__":
    unittest.main()
