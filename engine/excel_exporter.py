import logging
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

class ExcelExporter:
    """
    用于将 AI 提取的润色建议导出为并排对比的 Excel 表格。
    支持用颜色标出接受（绿色）或驳回（黄色）的理由，方便复审。
    """
    def __init__(self):
        self.logger = logging.getLogger("ExcelExporter")
        self.HEADER = ["段落号", "原句 (旧)", "润色后 (新)", "AI 修改理由", "状态"]
        self.COL_WIDTHS = [10, 50, 50, 30, 15]

        self.FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
        self.FILL_KEPT   = PatternFill("solid", fgColor="D4EDDA") # 绿色 (保留)
        self.FILL_REJ    = PatternFill("solid", fgColor="FFF3CD") # 黄色 (驳回)
        self.FONT_HEADER = Font(color="FFFFFF", bold=True)
        self.ALIGN_TOP   = Alignment(wrap_text=True, vertical="top")

    def export(self, export_path: str, records: list):
        """
        records 期待的数据格式:
        [
            {
               "paragraph_idx": 15,
               "old": "被广泛应用",
               "new": "应用广泛",
               "reason": "更加顺口",
               "status": "✅保留" # 或是 "❌驳回"
            }, ...
        ]
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "润色审核明细"

        # 写表头
        ws.append(self.HEADER)
        for cell in ws[1]:
            cell.fill = self.FILL_HEADER
            cell.font = self.FONT_HEADER
            cell.alignment = self.ALIGN_TOP

        # 写数据
        for item in records:
            row_data = [
                item.get("paragraph_idx", ""),
                item.get("old", ""),
                item.get("new", ""),
                item.get("reason", ""),
                item.get("status", "")
            ]
            ws.append(row_data)
            row_num = ws.max_row
            
            # 着色
            status = row_data[4]
            fill = None
            if status == "✅保留":
                fill = self.FILL_KEPT
            elif status == "❌驳回":
                fill = self.FILL_REJ
                
            for cell in ws[row_num]:
                if fill:
                    cell.fill = fill
                cell.alignment = self.ALIGN_TOP

        # 冻结首行
        ws.freeze_panes = "A2"

        # 列宽调整
        col_letters = "ABCDE"
        for i, width in enumerate(self.COL_WIDTHS):
            if i < len(col_letters):
                ws.column_dimensions[col_letters[i]].width = width

        try:
            wb.save(export_path)
            self.logger.info(f"成功导出 Excel 报表: {export_path}")
            return True
        except Exception as e:
            self.logger.error(f"导出 Excel 失败: {e}")
            return False
